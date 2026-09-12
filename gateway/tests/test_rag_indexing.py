"""app/rag.py の純粋関数(_chunk_text, _extract_text)と index_document のテスト。

推論バックエンドは conftest の autouse fixture で app.rag.get_embed_backend が
FakeBackend に差し替わる。実ファイルI/Oは tmp_path 上で行い、
ASYNC(ruff)対策として async テスト内では asyncio.to_thread 経由にする。
"""

import asyncio
from collections.abc import Awaitable, Callable
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document as DocxDocument
from sqlalchemy import select, update
from sqlalchemy.ext.asyncio import AsyncSession

from app.config import get_settings
from app.models import Chunk, Document, Model
from app.rag import _chunk_text, _extract_text, index_document
from tests.conftest import FakeBackend

_MINIMAL_PDF_CONTENT = b"BT /F1 24 Tf 10 100 Td (Hello RAG) Tj ET"
_MINIMAL_PDF = (
    b"%PDF-1.4\n"
    b"1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj\n"
    b"2 0 obj<</Type/Pages/Kids[3 0 R]/Count 1>>endobj\n"
    b"3 0 obj<</Type/Page/Parent 2 0 R/Resources<</Font<</F1 4 0 R>>>>"
    b"/MediaBox[0 0 200 200]/Contents 5 0 R>>endobj\n"
    b"4 0 obj<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>endobj\n"
    b"5 0 obj<</Length "
    + str(len(_MINIMAL_PDF_CONTENT)).encode()
    + b">>\nstream\n"
    + _MINIMAL_PDF_CONTENT
    + b"\nendstream\nendobj\n"
    b"xref\n0 6\n0000000000 65535 f \n"
    b"trailer<</Size 6/Root 1 0 R>>\n"
    b"startxref\n0\n"
    b"%%EOF"
)


async def _write_bytes(path: Path, data: bytes) -> None:
    await asyncio.to_thread(path.write_bytes, data)


# ---------------------------------------------------------------------------
# _chunk_text: 固定長+オーバーラップの境界(literal assertion)
# ---------------------------------------------------------------------------


def test_chunk_text_empty_returns_no_chunks() -> None:
    assert _chunk_text("", chunk_size=10, overlap=3) == []


def test_chunk_text_whitespace_only_returns_no_chunks() -> None:
    assert _chunk_text("   \n\t  ", chunk_size=10, overlap=3) == []


def test_chunk_text_shorter_than_chunk_size_returns_single_chunk() -> None:
    assert _chunk_text("short", chunk_size=10, overlap=3) == ["short"]


def test_chunk_text_overlap_boundary_literal() -> None:
    text = "0123456789ABCDEFGHIJ"  # 20文字, chunk_size=10, overlap=3 -> step=7
    result = _chunk_text(text, chunk_size=10, overlap=3)
    assert result == ["0123456789", "789ABCDEFG", "EFGHIJ"]
    # 隣接チャンクが overlap 分だけ重なっていること
    assert result[0][-3:] == result[1][:3]
    assert result[1][-3:] == result[2][:3]


def test_chunk_text_overlap_greater_equal_chunk_size_steps_by_one() -> None:
    """overlap >= chunk_size のとき、step = max(1, chunk_size - overlap) が

    1にクランプされ、無限ループにならず1文字刻みで進むこと。"""
    result = _chunk_text("abcde", chunk_size=2, overlap=5)
    assert result == ["ab", "bc", "cd", "de"]


# ---------------------------------------------------------------------------
# _extract_text: 拡張子ごとの実抽出(PDF/DOCXは実際にライブラリでパースする)
# ---------------------------------------------------------------------------


def test_extract_text_txt(tmp_path: Path) -> None:
    path = tmp_path / "a.txt"
    path.write_text("日本語のテキストです", encoding="utf-8")
    assert _extract_text(path, ".txt") == "日本語のテキストです"


def test_extract_text_pdf(tmp_path: Path) -> None:
    path = tmp_path / "a.pdf"
    path.write_bytes(_MINIMAL_PDF)
    assert "Hello RAG" in _extract_text(path, ".pdf")


def test_extract_text_docx(tmp_path: Path) -> None:
    doc = DocxDocument()
    doc.add_paragraph("Hello RAG docx")
    buf = BytesIO()
    doc.save(buf)
    path = tmp_path / "a.docx"
    path.write_bytes(buf.getvalue())
    assert "Hello RAG docx" in _extract_text(path, ".docx")


def test_extract_text_unsupported_extension_raises(tmp_path: Path) -> None:
    path = tmp_path / "a.exe"
    path.write_bytes(b"\x00")
    with pytest.raises(ValueError, match="未対応"):
        _extract_text(path, ".exe")


# ---------------------------------------------------------------------------
# index_document: pending -> indexing -> ready / 失敗時 failed
# ---------------------------------------------------------------------------


async def _make_pending_document(db: AsyncSession, owner_id: int, storage_path: Path) -> Document:
    document = Document(
        owner_id=owner_id,
        filename=f"pytest-index-{storage_path.name}",
        mime_type="text/plain",
        size_bytes=0,  # index_document はこのフィールドを参照しないため未使用
        storage_path=str(storage_path),
        status="pending",
    )
    db.add(document)
    await db.commit()
    await db.refresh(document)
    return document


async def test_index_document_happy_path_creates_embedded_chunks(
    db: AsyncSession,
    login_as_new_user: Callable,
    make_model: Callable,
    isolate_model: Callable[[Model], Awaitable[None]],
    tmp_path: Path,
    fake_backend: FakeBackend,
) -> None:
    user = await login_as_new_user()
    # id昇順で最初の有効モデルが選ばれるため、isolateしないとseedの
    # embed-standardが使われてしまい、このテストが作ったモデルが実際には
    # 使われないまま緑になる(過去そうだった)。
    embed_model = await make_model(kind="embedding")
    await isolate_model(embed_model)

    path = tmp_path / "note.txt"
    await _write_bytes(path, b"0123456789ABCDEFGHIJ")
    document = await _make_pending_document(db, user.id, path)

    await index_document(document.id)

    await db.refresh(document)
    assert document.status == "ready"

    result = await db.execute(
        select(Chunk).where(Chunk.document_id == document.id).order_by(Chunk.ordinal)
    )
    chunks = result.scalars().all()
    assert len(chunks) == 1  # 既定のchunk_size(1000)なので短文は1チャンクに収まる
    assert chunks[0].content == "0123456789ABCDEFGHIJ"
    assert len(chunks[0].embedding) == 1024
    # このテストが作った embed_model が実際に使われたこと(値・引数とも)。
    assert fake_backend.embed_calls == [(["0123456789ABCDEFGHIJ"], embed_model.backend_name)]


async def test_index_document_missing_file_marks_failed(
    db: AsyncSession,
    login_as_new_user: Callable,
    make_model: Callable,
    tmp_path: Path,
) -> None:
    user = await login_as_new_user()
    await make_model(kind="embedding")

    missing_path = tmp_path / "does-not-exist.txt"
    document = await _make_pending_document(db, user.id, missing_path)

    await index_document(document.id)  # 例外を外に投げないこと

    await db.refresh(document)
    assert document.status == "failed"
    result = await db.execute(select(Chunk).where(Chunk.document_id == document.id))
    assert result.scalars().all() == []


async def test_index_document_empty_text_marks_failed(
    db: AsyncSession,
    login_as_new_user: Callable,
    make_model: Callable,
    tmp_path: Path,
) -> None:
    user = await login_as_new_user()
    await make_model(kind="embedding")

    path = tmp_path / "empty.txt"
    await _write_bytes(path, b"   \n  ")
    document = await _make_pending_document(db, user.id, path)

    await index_document(document.id)

    await db.refresh(document)
    assert document.status == "failed"


async def test_index_document_no_embedding_model_marks_failed(
    db: AsyncSession,
    login_as_new_user: Callable,
    tmp_path: Path,
) -> None:
    user = await login_as_new_user()

    # 有効なembeddingモデルを全て一時的に無効化する(seedデータ含む)。
    result = await db.execute(
        select(Model.id).where(Model.kind == "embedding", Model.is_enabled.is_(True))
    )
    disabled_ids = [row[0] for row in result.all()]
    if disabled_ids:
        await db.execute(update(Model).where(Model.id.in_(disabled_ids)).values(is_enabled=False))
        await db.commit()

    try:
        path = tmp_path / "note.txt"
        await _write_bytes(path, b"hello")
        document = await _make_pending_document(db, user.id, path)

        await index_document(document.id)

        await db.refresh(document)
        assert document.status == "failed"
    finally:
        if disabled_ids:
            await db.execute(
                update(Model).where(Model.id.in_(disabled_ids)).values(is_enabled=True)
            )
            await db.commit()


async def test_index_document_is_idempotent_on_rerun(
    db: AsyncSession,
    login_as_new_user: Callable,
    make_model: Callable,
    isolate_model: Callable[[Model], Awaitable[None]],
    tmp_path: Path,
) -> None:
    """再実行してもchunksが重複せず、常に最新の抽出結果で置き換わること。"""
    user = await login_as_new_user()
    embed_model = await make_model(kind="embedding")
    await isolate_model(embed_model)

    path = tmp_path / "note.txt"
    await _write_bytes(path, b"0123456789ABCDEFGHIJ")
    document = await _make_pending_document(db, user.id, path)

    await index_document(document.id)
    await index_document(document.id)

    result = await db.execute(select(Chunk).where(Chunk.document_id == document.id))
    chunks = result.scalars().all()
    assert len(chunks) == 1  # 2回実行しても1件のまま(delete→insertで置き換わる)


class IndexedEmbedBackend(FakeBackend):
    """i番目のテキストに対し、index i成分だけ1.0の見分けやすいベクトルを返す。

    ordinal・content・embeddingの対応がずれていないかを検証するための道具。
    """

    async def embed(self, texts: list[str], model: str) -> list[list[float]]:
        self.embed_calls.append((texts, model))
        vectors = []
        for i in range(len(texts)):
            v = [0.0] * 1024
            v[i] = 1.0
            vectors.append(v)
        return vectors


async def test_index_document_multi_chunk_ordinal_content_embedding_correspondence(
    db: AsyncSession,
    login_as_new_user: Callable,
    make_model: Callable,
    isolate_model: Callable[[Model], Awaitable[None]],
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """複数チャンクがある場合に、ordinal・content・embeddingの対応がずれないこと。

    確認済みの意図: app/rag.py の `zip(chunks, embeddings, strict=True)` の
    対応を(例えば embeddings 側だけ逆順にするなどして)ずらすと、本テストの
    embedding 側アサーションが失敗する。1チャンクのテストだけでは
    この種の対応ずれ(検索結果が静かに壊れる、最も致命的なバグ形)を検知できない。
    """
    user = await login_as_new_user()
    embed_model = await make_model(kind="embedding")
    await isolate_model(embed_model)

    settings = get_settings()
    monkeypatch.setattr(settings, "chunk_size_chars", 10)
    monkeypatch.setattr(settings, "chunk_overlap_chars", 3)

    path = tmp_path / "note.txt"
    await _write_bytes(path, b"0123456789ABCDEFGHIJ")
    document = await _make_pending_document(db, user.id, path)

    fake = IndexedEmbedBackend()
    monkeypatch.setattr("app.rag.get_embed_backend", lambda: fake)

    await index_document(document.id)

    await db.refresh(document)
    assert document.status == "ready"

    result = await db.execute(
        select(Chunk).where(Chunk.document_id == document.id).order_by(Chunk.ordinal)
    )
    chunks = result.scalars().all()
    assert [c.content for c in chunks] == ["0123456789", "789ABCDEFG", "EFGHIJ"]
    for i, chunk in enumerate(chunks):
        assert chunk.ordinal == i
        assert chunk.embedding[i] == pytest.approx(1.0)
        others = [chunk.embedding[j] for j in range(len(chunks)) if j != i]
        assert all(v == pytest.approx(0.0) for v in others)

    assert fake.embed_calls[0][1] == embed_model.backend_name


async def test_index_document_nonexistent_id_is_noop(db: AsyncSession) -> None:
    """存在しない document_id を渡しても例外を投げず、早期returnすること。"""
    result_before = await db.execute(select(Chunk))
    count_before = len(result_before.scalars().all())

    await index_document(999_999_999)

    result_after = await db.execute(select(Chunk))
    assert len(result_after.scalars().all()) == count_before
